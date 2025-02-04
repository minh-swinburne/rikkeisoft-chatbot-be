from sqlalchemy.ext.asyncio import AsyncSession
from app.repos.document import DocumentRepository
from app.core.settings import settings
from app.schemas import DocumentBase, DocumentModel, DocumentUpdate
from app.aws import s3
from typing import Optional
import io
import os


class DocumentService:
    """
    Handles business logic for document management, including text extraction, CRUD operations, and processing.
    """

    @staticmethod
    async def create_document(
        db: AsyncSession, doc_data: DocumentBase
    ) -> DocumentModel:
        """Create a new document in the database."""
        document = await DocumentRepository.create(db, doc_data)
        return DocumentModel.model_validate(document)

    @staticmethod
    async def list_documents(db: AsyncSession) -> list[DocumentModel]:
        """List all documents in the database."""
        docs = await DocumentRepository.list(db)
        # print("Docs:", docs[0].__dict__)
        return [DocumentModel.model_validate(doc) for doc in docs]

    @staticmethod
    async def get_document_by_id(
        db: AsyncSession, doc_id: str
    ) -> Optional[DocumentModel]:
        """Retrieve a document by its ID."""
        document = await DocumentRepository.get_by_id(db, doc_id)
        if not document:
            return None
        return DocumentModel.model_validate(document)

    @staticmethod
    async def update_document(
        db: AsyncSession, doc_id: str, updates: DocumentUpdate
    ) -> Optional[DocumentModel]:
        """Update document details."""
        document = await DocumentRepository.update(db, doc_id, updates)
        if not document:
            return None
        return DocumentModel.model_validate(document)

    @staticmethod
    async def delete_document(db: AsyncSession, doc_id: str) -> bool:
        """Delete a document by its ID."""
        from app.bot import vector_db

        document = await DocumentRepository.get_by_id(db, doc_id)
        object_name = os.path.join(
            settings.upload_folder, f"{doc_id}.{document.file_type}"
        )
        s3.delete_file(object_name)
        vector_db.delete_data(doc_id)
        return await DocumentRepository.delete(db, doc_id)

    @staticmethod
    async def generate_document_url(db: AsyncSession, doc_id: str) -> str:
        """Get the URL of a document for previewing."""
        document = await DocumentRepository.get_by_id(db, doc_id)
        object_name = os.path.join(
            settings.upload_folder, f"{doc_id}.{document.file_type}"
        )
        return s3.generate_presigned_url(object_name)

    @staticmethod
    async def upload_document(document: DocumentModel, file_content: bytes) -> str:
        """Upload file directly to S3 and return its key (path)."""
        file_obj = io.BytesIO(file_content)
        object_name = os.path.join(
            settings.upload_folder, f"{document.id}.{document.file_type}"
        )
        return s3.upload_file(object_name, file_obj, document.file_name)

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
        words = text.split()
        chunks = [
            " ".join(words[i : i + chunk_size])
            for i in range(0, len(words), chunk_size - overlap)
        ]
        return chunks

    @staticmethod
    def extract_text(file_content: bytes, file_type: str) -> str:
        """Extract text from a file based on its type."""
        extractors = {
            "pdf": DocumentService._extract_text_from_pdf,
            "docx": DocumentService._extract_text_from_doc,
            "xlsx": DocumentService._extract_text_from_excel,
            "html": DocumentService._extract_text_from_html,
        }
        if file_type not in extractors:
            raise ValueError(f"Unsupported file type: {file_type}")
        return extractors[file_type](file_content)

    @staticmethod
    async def embed_document(document: DocumentModel, file_content: bytes) -> bool:
        from app.bot.embedding import get_embedding
        from app.bot.vector_db import insert_data

        text = DocumentService.extract_text(file_content, document.file_type)
        chunks = DocumentService.chunk_text(text)
        embeddings = get_embedding(chunks)
        data = []

        for chunk, embedding in zip(chunks, embeddings):
            data.append(
                {
                    "document_id": document.id,
                    "text": chunk,
                    "embedding": embedding,
                }
            )

        insert_data(data)

    @staticmethod
    def _extract_text_from_pdf(file_content: bytes) -> str:
        from PIL import Image
        import pytesseract
        import fitz  # PyMuPDF for PDF

        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        text = ""
        with fitz.open(stream=file_content, filetype="pdf") as pdf:
            for page in pdf:
                page_text = page.get_text()
                if not page_text.strip():
                    images = page.get_images(full=True)
                    for img_index, img in enumerate(images):
                        xref = img[0]
                        base_image = pdf.extract_image(xref)
                        image_bytes = base_image["image"]
                        image = Image.open(io.BytesIO(image_bytes))
                        ocr_text = pytesseract.image_to_string(image)
                        text += f"OCR - Page {page.number}, Image {img_index + 1}:\n{ocr_text}\n"
                else:
                    text += page_text
        return text

    @staticmethod
    def _extract_text_from_doc(file_content: bytes) -> str:
        from docx import Document as Docx  # For .docx

        file_obj = io.BytesIO(file_content)
        doc = Docx(file_obj)
        return "\n".join([p.text for p in doc.paragraphs])

    @staticmethod
    def _extract_text_from_excel(file_content: bytes) -> str:
        import openpyxl  # For .xlsx

        file_obj = io.BytesIO(file_content)
        text = ""
        workbook = openpyxl.load_workbook(file_obj)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join(map(str, row)) + "\n"
        return text

    @staticmethod
    def _extract_text_from_html(file_content: bytes) -> str:
        from bs4 import BeautifulSoup  # For .html

        html_content = file_content.decode("utf-8")
        soup = BeautifulSoup(html_content, "html.parser")
        return soup.get_text()
