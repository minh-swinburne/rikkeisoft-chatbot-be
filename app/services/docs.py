from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from app.core.config import settings
from app.models.documents import DocumentBase
from typing import List, Optional
from bs4 import BeautifulSoup  # For .html
from docx import Document   # For .docx
from PIL import Image
import pytesseract
import openpyxl  # For .xlsx
import uuid
import fitz  # PyMuPDF for PDF
import io




pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

# print(settings.tesseract_cmd)


# Create a new document entry
async def create_document(
    db: AsyncSession,
    filename: str,
    file_type: str,
    title: str,
    description: str,
    categories: list[str],
    creator: str,
    created_date,
    restricted: bool,
    uploader: str,
):
    doc = DocumentBase(
        id=str(uuid.uuid4()),
        filename=filename,
        file_type=file_type,
        title=title,
        description=description,
        categories=",".join(categories),  # Convert list to a comma-separated string
        creator=creator,
        created_date=created_date,
        restricted=restricted,
        uploader=uploader,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return doc


# Retrieve a document by ID
async def get_document_by_id(db: AsyncSession, doc_id: str):
    result = await db.execute(select(DocumentBase).where(DocumentBase.id == doc_id))
    print("Fetching document by ID:", result)
    return result.scalar_one_or_none()


def extract_text(file_path: str, file_type: str) -> str:
    """
    Extract text content from a given file based on its type.
    """
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "docx":
        return extract_text_from_doc(file_path)
    elif file_type == "xlsx":
        return extract_text_from_excel(file_path)
    elif file_type == "web_content":
        return extract_text_from_html(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    with fitz.open(file_path) as pdf:
        for page in pdf:
            page_text = page.get_text()
            if not page_text.strip():  # Page might be an image
                images = page.get_images(full=True)
                for img_index, img in enumerate(images):
                    xref = img[0]
                    base_image = pdf.extract_image(xref)
                    image_bytes = base_image["image"]
                    image = Image.open(io.BytesIO(image_bytes))
                    ocr_text = pytesseract.image_to_string(image)
                    text += f"OCR - Page {page.number}, Image {img_index + 1}:\n{ocr_text}\n"
            else:
                text += page_text
    return text


def extract_text_from_doc(file_path: str) -> str:
    doc = Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])


def extract_text_from_excel(file_path: str) -> str:
    text = ""
    workbook = openpyxl.load_workbook(file_path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += " ".join(map(str, row)) + "\n"
    return text


def extract_text_from_html(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as file:
        soup = BeautifulSoup(file, "html.parser")
        return soup.get_text()


async def update_document(
    db: AsyncSession,
    doc_id: str,
    title: Optional[str] = None,
    description: Optional[str] = None,
    categories: Optional[List[str]] = None,
    restricted: Optional[bool] = None
):
    print("Updating document with ID: ", doc_id)
    print("New values - title:", title, "description:", description, "categories:", categories, "restricted:", restricted)

    # Retrieve the document by ID
    document = await get_document_by_id(db, doc_id)
    if not document:
        print("Document not found for ID: ", doc_id)
        return None

    print("Document retrieved: ", document)

    # Update fields only if new values are provided
    print(f"Before update: {document}")
    if title:
        document.title = title
    if description:
        document.description = description
    if categories:
        document.categories = ",".join(categories)
    if restricted is not None:
        document.restricted = restricted

    print(f"After update: {document}")


    # Add and commit the updated document
    db.add(document)
    await db.commit()
    await db.refresh(document)

    print("Document updated successfully: ", document)
    return document

async def delete_document(db: AsyncSession, doc_id: str):

    document = await get_document_by_id(db, doc_id)
    if not document:
        return False
    await db.delete(document)
    await db.commit()
    return True
